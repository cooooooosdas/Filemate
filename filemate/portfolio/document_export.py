"""只排版已保存事实和报告，不调用模型或加载远程资源。"""

from __future__ import annotations

import io
import re
import threading
from html import escape
from pathlib import Path

_FONT_LOCK = threading.Lock()


def _lines(markdown: str) -> list[tuple[int, str]]:
    if len(markdown.encode('utf-8')) > 2 * 1024 * 1024:
        raise ValueError('文档内容超过2MB，请缩小报告期间或使用Markdown/JSON导出')
    lines = markdown.splitlines()
    if len(lines) > 20000:
        raise ValueError('文档段落过多，请缩小报告期间或使用Markdown/JSON导出')
    result = []
    for line in lines:
        line = ''.join(char for char in line if ord(char) >= 32 or char == '\t')
        heading = re.match(r'^(#{1,3})\s+(.*)$', line)
        result.append((len(heading[1]) if heading else 0, heading[2] if heading else line))
    return result


def document_pdf(markdown: str, *, title: str) -> bytes:
    """嵌入已有开放许可中文字体，自动分页并保留可选择正文。"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    lines, font = _lines(markdown), 'FileMatePortfolioNotoSC'
    with _FONT_LOCK:
        if font not in pdfmetrics.getRegisteredFontNames():
            font_path = Path(__file__).parents[1] / 'interview_review' / 'assets' / 'NotoSansSC-Regular.ttf'
            pdfmetrics.registerFont(TTFont(font, str(font_path)))
    body = ParagraphStyle('portfolio-body', fontName=font, fontSize=10.5, leading=17,
                          wordWrap='CJK', textColor=colors.HexColor('#15264A'), spaceAfter=7)
    styles = {0: body}
    for level, size in ((1, 24), (2, 16), (3, 12)):
        styles[level] = ParagraphStyle(f'portfolio-h{level}', parent=body, fontSize=size,
                                      leading=size * 1.45, spaceBefore=12 if level > 1 else 0,
                                      spaceAfter=9, keepWithNext=True)
    story = [Paragraph(escape(value).replace('\t', '    '), styles[level]) if value else Spacer(1, 4)
             for level, value in lines]
    if not story:
        story.append(Paragraph('暂无内容', body))

    def footer(canvas, document):
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor('#405573'))
        canvas.drawString(42, 24, 'FileMate')
        canvas.drawRightString(A4[0] - 42, 24, str(document.page))

    stream = io.BytesIO()
    SimpleDocTemplate(stream, pagesize=A4, leftMargin=42, rightMargin=42,
                      topMargin=42, bottomMargin=42, title=title, author='FileMate').build(
                          story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()


def document_docx(markdown: str, *, title: str) -> bytes:
    """创建可编辑Word正文，用户文字不作为XML或嵌入对象执行。"""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    lines, document = _lines(markdown), Document()
    document.core_properties.title = title
    document.core_properties.author = 'FileMate'
    for section in document.sections:
        section.top_margin = section.bottom_margin = Inches(0.65)
        section.left_margin = section.right_margin = Inches(0.65)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        style = document.styles[name]
        style.font.name = 'Microsoft YaHei'
        style.font.color.rgb = RGBColor.from_string('15264A')
        properties = style.element.get_or_add_rPr()
        fonts = properties.find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts')
            properties.append(fonts)
        fonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    document.styles['Normal'].font.size = Pt(10.5)
    document.styles['Normal'].paragraph_format.line_spacing = 1.45
    for level, value in lines:
        if level:
            document.add_heading(value, level=level)
        else:
            document.add_paragraph(value)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()
