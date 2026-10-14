"""Regression coverage for document resource exhaustion and event-loop starvation."""

from __future__ import annotations

import asyncio
import os
import subprocess
import sys
import time
from contextvars import ContextVar
from pathlib import Path

import pytest

from filemate.perception import FileParser
from filemate.perception import bounded_parse as boundary


def test_real_office_and_pdf_enrichment(tmp_path: Path) -> None:
    from docx import Document
    from reportlab.pdfgen.canvas import Canvas

    doc = Document()
    doc.add_paragraph('Synthetic lecture: stacks are last-in first-out.')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Stack'
    table.cell(0, 1).text = 'LIFO'
    path = tmp_path / 'lecture.docx'
    doc.save(path)
    result = FileParser().parse_enriched(path)
    assert not result.get('error'), result.get('error')
    assert 'LIFO' in result['raw_text'] and result['metadata']['tables'] == 1
    path = tmp_path / 'lecture.pdf'
    canvas = Canvas(str(path))
    for index in range(30):
        canvas.drawString(40, 700, f'Synthetic page {index}: stack lecture')
        canvas.showPage()
    canvas.save()
    result = FileParser().parse_enriched(path)
    assert not result.get('error'), result.get('error')
    assert result['metadata']['text_pages'] == 30


def test_long_pdf_preserves_text_and_marks_optional_structure_limit(tmp_path):
    from reportlab.pdfgen.canvas import Canvas

    path = tmp_path / 'long.pdf'
    canvas = Canvas(str(path))
    for _ in range(51):
        canvas.drawString(40, 700, 'Synthetic lecture')
        canvas.showPage()
    canvas.save()
    result = FileParser().parse_enriched(path)
    assert not result.get('error')
    assert result['metadata']['text_pages'] == 51
    assert result['metadata']['enrichment_incomplete'] is True
    assert result['note']


def test_large_page_count_rejected_before_ai(tmp_path: Path) -> None:
    from PyPDF2 import PdfWriter

    path = tmp_path / 'oversized.pdf'
    writer = PdfWriter()
    for _ in range(501):
        writer.add_blank_page(width=100, height=100)
    with path.open('wb') as output:
        writer.write(output)
    result = FileParser().parse_enriched(path)
    assert '500' in result['error'] and not result['raw_text']


@pytest.mark.parametrize('reason', ['timeout', 'memory'])
def test_worker_is_killed_reaped_and_slot_reusable(tmp_path, monkeypatch, reason):
    original = subprocess.Popen
    processes = []
    monkeypatch.setenv('LLM_API_KEY', 'private-test-credential')

    def launch(command, **kwargs):
        assert not any('lecture' in arg for arg in command)
        assert 'LLM_API_KEY' not in kwargs['env']
        body = 'import sys,time;sys.stdin.buffer.read();time.sleep(60)'
        process = original([sys.executable, '-c', body], **kwargs)
        processes.append(process)
        return process

    monkeypatch.setattr(boundary.subprocess, 'Popen', launch)
    monkeypatch.setattr(boundary, 'WALL_SECONDS', 0.15 if reason == 'timeout' else 5)
    if reason == 'memory':
        monkeypatch.setattr(boundary, 'MEMORY_BYTES', 1024)
    path = tmp_path / 'lecture.pdf'
    path.write_bytes(b'bad')
    start = time.monotonic()
    result = FileParser().parse(path)
    assert result['error'] and time.monotonic() - start < 6
    assert processes[0].poll() is not None
    assert boundary._SLOT.acquire(blocking=False)
    boundary._SLOT.release()


def test_busy_worker_returns_without_creating_another(tmp_path):
    path = tmp_path / 'lecture.pdf'
    path.write_bytes(b'bad')
    assert boundary._SLOT.acquire(blocking=False)
    try:
        result = FileParser().parse(path)
        assert '稍后' in result['error']
    finally:
        boundary._SLOT.release()


@pytest.mark.skipif(sys.platform != 'linux', reason='Linux worker address-space limit')
def test_compressed_pdf_expansion_cannot_kill_parent(tmp_path):
    from PyPDF2 import PdfWriter
    from PyPDF2.generic import DecodedStreamObject, NameObject

    path = tmp_path / 'expansion.pdf'
    writer = PdfWriter()
    page = writer.add_blank_page(width=100, height=100)
    stream = DecodedStreamObject()
    stream.set_data(b' ' * (160 * 1024 * 1024))
    page[NameObject('/Contents')] = writer._add_object(stream.flate_encode())
    with path.open('wb') as output:
        writer.write(output)
    del stream, writer, page
    result = FileParser().parse(path)
    assert result.get('error')
    assert os.getpid() > 0


def test_processing_stage_keeps_loop_responsive_and_request_context(tmp_path, monkeypatch):
    import main

    scope = ContextVar('synthetic_request_scope', default='missing')
    monkeypatch.setenv('LLM_API_KEY', 'sk-test')
    entered = asyncio.Event()

    def stage(session):
        assert scope.get() == 'tenant-test'
        time.sleep(0.2)
        return session

    monkeypatch.setattr(main, '_make_stages', lambda *a, **k: [stage])

    async def scenario():
        scope.set('tenant-test')
        task = asyncio.create_task(main.process_single('synthetic.txt', db_path=str(tmp_path / 'x.db')))
        await asyncio.sleep(0.05)
        assert not task.done()
        entered.set()
        await task

    asyncio.run(scenario())
    assert entered.is_set()
